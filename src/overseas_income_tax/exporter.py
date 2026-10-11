from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import CalculationRequest, CalculationResult


HEADER = PatternFill("solid", fgColor="1F4E78")
SUBHEADER = PatternFill("solid", fgColor="D9EAF7")


def _header(row) -> None:
    for cell in row:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(horizontal="center")


def build_xlsx(request: CalculationRequest, result: CalculationResult) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "范围内测算"
    summary.append(["境外收入范围内测算", "值"])
    _header(summary[1])
    rows = [
        ("纳税年度", result.tax_year),
        ("结果状态", result.status.value),
        ("规则快照", result.rule_snapshot_version),
        ("计算指纹", result.fingerprint),
        ("范围内应纳税所得额（人民币）", float(result.taxable_income_cny)),
        ("范围内应纳税额（人民币）", float(result.tax_before_credit_cny)),
        ("本次采用的境外抵免（人民币）", float(result.credit_used_cny)),
        ("范围内预计补税（人民币）", None if result.range_supplemental_tax_cny is None else float(result.range_supplemental_tax_cny)),
    ]
    for row in rows:
        summary.append(row)
    summary.append([])
    summary.append(["提示", "存在范围外或待核验事项时，本表不得称为全年最终应补缴总额。"])

    details = workbook.create_sheet("已核验经济事项")
    details.append(["事项ID", "类别", "国家/地区", "日期", "收入币种", "税款币种", "税前收入", "报表扣税", "符合条件已缴税", "证据ID", "确认原因"])
    _header(details[1])
    included = set(result.included_event_ids)
    for event in request.events:
        if event.id not in included:
            continue
        details.append([
            event.id,
            event.category.value,
            event.country_or_region,
            event.occurred_on.isoformat() if event.occurred_on else "",
            event.currency,
            event.foreign_tax_currency,
            float(event.gross_amount),
            float(event.foreign_tax_reported or 0),
            float(event.foreign_tax_eligible or 0),
            ", ".join(event.evidence_ids),
            event.review_reason,
        ])

    credits = workbook.create_sheet("分国抵免账")
    credits.append(["国家/地区", "应纳税所得额", "抵免限额", "报表扣税额", "符合条件已缴税额", "本年采用抵免额", "未采用超限额"])
    _header(credits[1])
    for account in result.country_credit_accounts:
        credits.append([
            account.country_or_region,
            float(account.taxable_income_cny),
            float(account.credit_limit_cny),
            float(account.reported_foreign_tax_cny),
            float(account.eligible_foreign_tax_cny),
            float(account.credit_used_cny),
            float(account.excess_not_used_cny),
        ])

    derivations = workbook.create_sheet("计算推导")
    derivations.append([
        "事项ID", "来源文件", "页码", "证据位置", "原文片段", "原币收入",
        "收入币种", "折算汇率", "汇率日期", "汇率来源", "折合人民币",
        "报表扣税", "符合条件已缴税", "税款币种", "税款汇率", "推导公式",
    ])
    _header(derivations[1])
    for item in result.event_derivations:
        derivations.append([
            item.event_id,
            item.file_name,
            item.page,
            item.locator,
            item.excerpt,
            float(item.gross_amount),
            item.currency,
            float(item.cny_per_unit),
            item.rate_date.isoformat(),
            item.rate_source,
            float(item.gross_cny),
            float(item.foreign_tax_reported),
            float(item.foreign_tax_eligible),
            item.foreign_tax_currency,
            float(item.foreign_tax_cny_per_unit),
            item.formula,
        ])
    derivations.append([])
    derivations.append(["汇总项目", "结果（人民币）", "推导公式"])
    _header(derivations[derivations.max_row])
    for item in result.summary_derivations:
        derivations.append([item.label, float(item.result_cny), item.formula])

    audit = workbook.create_sheet("待核验与规则")
    audit.append(["阻断", "代码", "事项ID", "说明"])
    _header(audit[1])
    for item in result.pending_items:
        audit.append(["是" if item.blocking else "否", item.code, item.event_id or "", item.message])
    if request.audit_reviews:
        audit.append([])
        audit.append(["抽检事项ID", "结果", "备注"])
        _header(audit[audit.max_row])
        for review in request.audit_reviews:
            audit.append([
                review.event_id,
                "一致" if review.status.value == "consistent" else "不一致",
                review.note,
            ])
    audit.append([])
    audit.append(["汇率币种", "汇率日期", "1原币=CNY", "报价单位", "规则编号", "来源"])
    _header(audit[audit.max_row])
    for rate in request.exchange_rates:
        audit.append([rate.currency, rate.rate_date.isoformat(), float(rate.cny_per_unit), rate.quote_unit, rate.rule_code, rate.source])
    audit.append([])
    audit.append(["官方依据"])
    audit[audit.max_row][0].fill = SUBHEADER
    audit[audit.max_row][0].font = Font(bold=True)
    for url in request.rule_snapshot.source_urls:
        audit.append([url])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            letter = column[0].column_letter
            max_length = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 46)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
