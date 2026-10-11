from io import BytesIO
from dataclasses import replace

from openpyxl import load_workbook

from overseas_income_tax.engine import calculate
from overseas_income_tax.exporter import build_xlsx
from overseas_income_tax.models import AuditReview, AuditStatus
from test_engine import event, request


def test_excel_uses_same_calculation_result():
    calculation_request = request(event("one"))
    result = calculate(calculation_request)
    workbook = load_workbook(BytesIO(build_xlsx(calculation_request, result)), data_only=True)
    summary = workbook["范围内测算"]
    values = {row[0].value: row[1].value for row in summary.iter_rows() if row[0].value}
    assert values["范围内应纳税所得额（人民币）"] == float(result.taxable_income_cny)
    assert values["范围内预计补税（人民币）"] == float(result.range_supplemental_tax_cny)
    assert values["计算指纹"] == result.fingerprint
    derivation = workbook["计算推导"]
    assert derivation["A2"].value == "one"
    assert "1000 USD × 7" in derivation["P2"].value


def test_excel_records_random_audit_result_without_changing_calculation():
    calculation_request = request(event("one"))
    calculation_request = replace(
        calculation_request,
        audit_reviews=(AuditReview("one", AuditStatus.INCONSISTENT, "原件金额不同"),),
    )
    result = calculate(calculation_request)
    workbook = load_workbook(BytesIO(build_xlsx(calculation_request, result)), data_only=True)
    rows = list(workbook["待核验与规则"].values)
    assert ("one", "不一致", "原件金额不同", None, None, None) in rows
