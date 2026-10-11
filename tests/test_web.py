from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfWriter

from overseas_income_tax.models import to_primitive
from overseas_income_tax.web import create_app
from test_engine import event, request


def test_health_declares_privacy_defaults():
    client = create_app().test_client()
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json == {"status": "ok", "storage": "memory-only", "model": "disabled"}
    assert response.headers["X-Overseas-Income-Tax-API-Version"] == "v1"


def test_agent_can_discover_capabilities_and_request_schema():
    client = create_app().test_client()
    capability_response = client.get("/api/v1/capabilities")
    schema_response = client.get("/api/v1/schema/calculation-request")
    assert capability_response.status_code == 200
    assert capability_response.json["agent_vendor_neutral"] is True
    assert capability_response.json["interfaces"]["http"]["calculate"] == "/api/v1/calculate"
    assert capability_response.json["recognition"]["supported_templates"] == [
        "irs_1042_s", "bank_savings_statement", "ibkr_activity_statement"
    ]
    assert schema_response.status_code == 200
    assert schema_response.json["$schema"] == "https://json-schema.org/draft/2020-12/schema"
    assert set(schema_response.json["required"]) == {
        "questionnaire", "events", "observations", "exchange_rates"
    }


def test_reference_rates_are_versioned_official_records():
    client = create_app().test_client()
    response = client.get("/api/v1/reference-rates/2025")
    assert response.status_code == 200
    assert {item["currency"] for item in response.json["rates"]} == {"USD", "HKD"}
    assert all(item["rate_date"] == "2025-12-31" for item in response.json["rates"])
    assert all("safe.gov.cn" in item["source"] for item in response.json["rates"])


def test_calculation_endpoint():
    client = create_app().test_client()
    payload = to_primitive(request(event("one")))
    response = client.post("/api/v1/calculate", json=payload)
    assert response.status_code == 200
    assert response.json["range_supplemental_tax_cny"] == "700.00"
    assert response.json["event_derivations"][0]["formula"] == "1000 USD × 7 CNY/USD = 7000.00 CNY"
    assert response.headers["Cache-Control"] == "no-store"


def test_api_rejects_string_boolean_instead_of_treating_it_as_true():
    client = create_app().test_client()
    payload = to_primitive(request(event("one")))
    payload["questionnaire"]["same_country_other_income_reviewed"] = "false"
    response = client.post("/api/v1/calculate", json=payload)
    assert response.status_code == 400
    assert "布尔值" in response.json["error"]


def test_api_rejects_non_finite_amount():
    client = create_app().test_client()
    payload = to_primitive(request(event("one")))
    payload["events"][0]["gross_amount"] = "NaN"
    response = client.post("/api/v1/calculate", json=payload)
    assert response.status_code == 400


def test_api_returns_json_for_a_malformed_body():
    client = create_app().test_client()
    response = client.post(
        "/api/v1/calculate",
        data="not-json",
        content_type="application/json",
    )
    assert response.status_code == 400
    assert "JSON 对象" in response.json["error"]


def test_client_cannot_override_the_server_rule_snapshot():
    client = create_app().test_client()
    payload = to_primitive(request(event("one")))
    payload["rule_snapshot"]["version"] = "agent-supplied-rule"
    payload["rule_snapshot"]["rate"] = "0.99"
    response = client.post("/api/v1/calculate", json=payload)
    assert response.status_code == 200
    assert response.json["rule_snapshot_version"] != "agent-supplied-rule"
    assert response.json["tax_before_credit_cny"] == "1400.00"


def test_pdf_extraction_is_request_scoped():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    content = BytesIO()
    writer.write(content)
    client = create_app().test_client()
    response = client.post(
        "/api/v1/extract",
        data={"file": (BytesIO(content.getvalue()), "sample.pdf")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    assert response.json["file_name"] == "sample.pdf"
    assert len(response.json["sha256"]) == 64
    assert response.json["pages"] == [{"page": 1, "text": ""}]
    assert response.json["recognition"]["status"] == "needs_manual_review"
    assert response.json["recognition"]["events"] == []
    assert "扫描件" in response.json["recognition"]["warnings"][0]


def test_export_endpoint_returns_same_result_workbook():
    client = create_app().test_client()
    payload = to_primitive(request(event("one")))
    response = client.post("/api/v1/export.xlsx", json=payload)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), data_only=True)
    assert workbook.sheetnames == ["范围内测算", "已核验经济事项", "分国抵免账", "计算推导", "待核验与规则"]


def test_legacy_endpoint_remains_a_compatible_alias():
    client = create_app().test_client()
    payload = to_primitive(request(event("one")))
    current = client.post("/api/v1/calculate", json=payload)
    legacy = client.post("/api/calculate", json=payload)
    assert legacy.status_code == 200
    assert legacy.json == current.json
